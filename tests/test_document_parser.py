from io import BytesIO

import pytest
from docx import Document as DocxDocument
from pypdf import PdfWriter
from pypdf.generic import DecodedStreamObject, DictionaryObject, NameObject

from app.rag.document_parser import DocumentParseError, parse_uploaded_document


def _parse(filename: str, data: bytes):
    return parse_uploaded_document(
        filename,
        data,
        max_characters=200_000,
        max_pdf_pages=300,
        max_docx_uncompressed_bytes=50 * 1024 * 1024,
    )


def _docx_bytes() -> bytes:
    document = DocxDocument()
    document.add_heading("差旅制度", level=1)
    document.add_paragraph("员工住宿标准为每晚 500 元，超出标准需要负责人审批。")
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "城市"
    table.cell(0, 1).text = "标准"
    table.cell(1, 0).text = "上海"
    table.cell(1, 1).text = "700 元"
    stream = BytesIO()
    document.save(stream)
    return stream.getvalue()


def _pdf_bytes() -> bytes:
    writer = PdfWriter()
    page = writer.add_blank_page(width=612, height=792)
    font = DictionaryObject(
        {
            NameObject("/Type"): NameObject("/Font"),
            NameObject("/Subtype"): NameObject("/Type1"),
            NameObject("/BaseFont"): NameObject("/Helvetica"),
        }
    )
    page[NameObject("/Resources")] = DictionaryObject(
        {
            NameObject("/Font"): DictionaryObject(
                {NameObject("/F1"): writer._add_object(font)}
            )
        }
    )
    content = DecodedStreamObject()
    content.set_data(
        b"BT /F1 12 Tf 72 720 Td (Refund period is eleven business days.) Tj ET"
    )
    page[NameObject("/Contents")] = writer._add_object(content)
    stream = BytesIO()
    writer.write(stream)
    return stream.getvalue()


def test_parse_docx_extracts_paragraphs_and_tables() -> None:
    parsed = _parse("差旅制度.docx", _docx_bytes())

    assert parsed.title == "差旅制度"
    assert "每晚 500 元" in parsed.content
    assert "上海 | 700 元" in parsed.content
    assert parsed.extension == ".docx"


def test_parse_real_pdf_extracts_text() -> None:
    parsed = _parse("refund-policy.pdf", _pdf_bytes())

    assert parsed.page_count == 1
    assert "Refund period is eleven business days" in parsed.content
    assert parsed.extension == ".pdf"


def test_parse_pdf_preserves_page_markers(monkeypatch) -> None:
    class FakePage:
        def __init__(self, text: str) -> None:
            self.text = text

        def extract_text(self) -> str:
            return self.text

    class FakeReader:
        is_encrypted = False
        pages = [FakePage("第一页的报销制度正文"), FakePage("第二页的审批流程正文")]

    monkeypatch.setattr("app.rag.document_parser.PdfReader", lambda *_args, **_kwargs: FakeReader())

    parsed = _parse("员工手册.pdf", b"%PDF-1.7 fake-test-content")

    assert parsed.page_count == 2
    assert "[第 1 页]" in parsed.content
    assert "第二页的审批流程正文" in parsed.content


def test_parse_pdf_rejects_image_only_content(monkeypatch) -> None:
    class EmptyPage:
        def extract_text(self) -> str:
            return ""

    class FakeReader:
        is_encrypted = False
        pages = [EmptyPage()]

    monkeypatch.setattr("app.rag.document_parser.PdfReader", lambda *_args, **_kwargs: FakeReader())

    with pytest.raises(DocumentParseError, match="OCR") as exc_info:
        _parse("扫描件.pdf", b"%PDF-1.7 fake-test-content")

    assert exc_info.value.code == "DOCUMENT_TEXT_EMPTY"


def test_parse_rejects_unsupported_file_type() -> None:
    with pytest.raises(DocumentParseError) as exc_info:
        _parse("notes.txt", b"plain text")

    assert exc_info.value.code == "DOCUMENT_TYPE_UNSUPPORTED"
