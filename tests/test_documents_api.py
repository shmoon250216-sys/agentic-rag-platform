from io import BytesIO

from docx import Document as DocxDocument
from fastapi.testclient import TestClient

from app.llm.client import llm_response_cache
from app.main import create_app
from app.rag.retriever import retrieval_cache


def test_documents_api_requires_token() -> None:
    client = TestClient(create_app())

    response = client.get("/api/v1/documents")

    assert response.status_code == 401


def test_documents_api_creates_and_lists_document() -> None:
    client = TestClient(create_app())
    headers = {"Authorization": "Bearer dev-token"}

    create_response = client.post(
        "/api/v1/documents",
        headers=headers,
        json={
            "title": "报销制度说明",
            "content": "这个系统使用 LangGraph 做 Supervisor 调度，并通过 RAG 检索企业制度内容。",
        },
    )
    list_response = client.get("/api/v1/documents", headers=headers)

    assert create_response.status_code == 200
    assert create_response.json()["document"]["title"] == "报销制度说明"
    assert list_response.status_code == 200
    assert list_response.json()["total_chunks"] >= 1


def test_document_changes_clear_hot_caches() -> None:
    client = TestClient(create_app())
    headers = {"Authorization": "Bearer dev-token"}
    retrieval_cache.set("demo", ["cached"])
    llm_response_cache.set("demo", "cached")

    response = client.post(
        "/api/v1/documents",
        headers=headers,
        json={
            "title": "缓存测试文档",
            "content": "新增文档会影响 RAG 检索结果，所以需要清理热点缓存。",
        },
    )

    assert response.status_code == 200
    assert retrieval_cache.stats()["size"] == 0
    assert llm_response_cache.stats()["size"] == 0


def test_documents_api_uploads_docx_and_indexes_extracted_text() -> None:
    client = TestClient(create_app())
    headers = {"Authorization": "Bearer dev-token"}
    document = DocxDocument()
    document.add_heading("产品手册", level=1)
    document.add_paragraph("星河系统的标准退款期限为七个工作日，特殊订单需要人工审核。")
    stream = BytesIO()
    document.save(stream)

    response = client.post(
        "/api/v1/documents/upload",
        headers=headers,
        data={"title": "退款业务手册"},
        files={
            "file": (
                "产品手册.docx",
                stream.getvalue(),
                "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            )
        },
    )

    assert response.status_code == 200
    assert response.json()["document"]["title"] == "退款业务手册"
    assert response.json()["document"]["chunk_count"] >= 1


def test_documents_api_rejects_unsupported_upload() -> None:
    client = TestClient(create_app())
    response = client.post(
        "/api/v1/documents/upload",
        headers={"Authorization": "Bearer dev-token"},
        files={"file": ("notes.txt", b"not a supported document", "text/plain")},
    )

    assert response.status_code == 400
    assert response.json()["error"]["code"] == "DOCUMENT_TYPE_UNSUPPORTED"
