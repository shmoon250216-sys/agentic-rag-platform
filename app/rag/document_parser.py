import re
from dataclasses import dataclass
from io import BytesIO
from pathlib import Path
from zipfile import BadZipFile, ZipFile

from docx import Document as DocxDocument
from docx.opc.exceptions import PackageNotFoundError
from pypdf import PdfReader
from pypdf.errors import FileNotDecryptedError, PdfReadError

SUPPORTED_DOCUMENT_EXTENSIONS = {".pdf", ".docx"}


@dataclass(frozen=True)
class ParsedDocument:
    title: str
    content: str
    extension: str
    page_count: int | None = None


class DocumentParseError(ValueError):
    def __init__(self, code: str, message: str) -> None:
        self.code = code
        self.message = message
        super().__init__(message)


def parse_uploaded_document(
    filename: str,
    data: bytes,
    *,
    max_characters: int,
    max_pdf_pages: int,
    max_docx_uncompressed_bytes: int,
) -> ParsedDocument:
    safe_name = Path(filename).name
    extension = Path(safe_name).suffix.lower()
    if extension not in SUPPORTED_DOCUMENT_EXTENSIONS:
        raise DocumentParseError(
            "DOCUMENT_TYPE_UNSUPPORTED",
            "仅支持 PDF 和 DOCX 文件",
        )
    if not data:
        raise DocumentParseError("DOCUMENT_EMPTY", "上传文件不能为空")

    if extension == ".pdf":
        content, page_count = _extract_pdf(data, max_pdf_pages=max_pdf_pages)
    else:
        content = _extract_docx(
            data,
            max_uncompressed_bytes=max_docx_uncompressed_bytes,
        )
        page_count = None

    normalized = _normalize_extracted_text(content)
    if len(normalized) < 20:
        raise DocumentParseError(
            "DOCUMENT_TEXT_EMPTY",
            "没有提取到足够的正文；扫描版 PDF 需要先进行 OCR",
        )
    if len(normalized) > max_characters:
        raise DocumentParseError(
            "DOCUMENT_TEXT_TOO_LARGE",
            f"提取后的正文超过 {max_characters} 个字符，请拆分文件后上传",
        )

    title = Path(safe_name).stem.strip()[:120] or "未命名文档"
    return ParsedDocument(
        title=title,
        content=normalized,
        extension=extension,
        page_count=page_count,
    )


def _extract_pdf(data: bytes, *, max_pdf_pages: int) -> tuple[str, int]:
    if not data.startswith(b"%PDF-"):
        raise DocumentParseError("PDF_INVALID", "文件内容不是有效的 PDF")

    try:
        reader = PdfReader(BytesIO(data), strict=False)
        if reader.is_encrypted:
            try:
                if reader.decrypt("") == 0:
                    raise DocumentParseError("PDF_ENCRYPTED", "暂不支持有密码的 PDF")
            except FileNotDecryptedError as exc:
                raise DocumentParseError("PDF_ENCRYPTED", "暂不支持有密码的 PDF") from exc

        page_count = len(reader.pages)
        if page_count > max_pdf_pages:
            raise DocumentParseError(
                "PDF_TOO_MANY_PAGES",
                f"PDF 超过 {max_pdf_pages} 页，请拆分后上传",
            )

        pages: list[str] = []
        for index, page in enumerate(reader.pages, start=1):
            text = page.extract_text() or ""
            if text.strip():
                pages.append(f"[第 {index} 页]\n{text.strip()}")
        return "\n\n".join(pages), page_count
    except DocumentParseError:
        raise
    except (PdfReadError, FileNotDecryptedError, ValueError, TypeError) as exc:
        raise DocumentParseError("PDF_INVALID", "PDF 已损坏或无法解析") from exc


def _extract_docx(data: bytes, *, max_uncompressed_bytes: int) -> str:
    _validate_docx_archive(data, max_uncompressed_bytes=max_uncompressed_bytes)
    try:
        document = DocxDocument(BytesIO(data))
    except (PackageNotFoundError, BadZipFile, KeyError, ValueError) as exc:
        raise DocumentParseError("DOCX_INVALID", "DOCX 已损坏或无法解析") from exc

    parts = [paragraph.text.strip() for paragraph in document.paragraphs if paragraph.text.strip()]
    for table in document.tables:
        rows = []
        for row in table.rows:
            cells = [re.sub(r"\s+", " ", cell.text).strip() for cell in row.cells]
            if any(cells):
                rows.append(" | ".join(cells))
        if rows:
            parts.append("\n".join(rows))
    return "\n\n".join(parts)


def _validate_docx_archive(data: bytes, *, max_uncompressed_bytes: int) -> None:
    if not data.startswith(b"PK"):
        raise DocumentParseError("DOCX_INVALID", "文件内容不是有效的 DOCX")
    try:
        with ZipFile(BytesIO(data)) as archive:
            names = set(archive.namelist())
            if "[Content_Types].xml" not in names or "word/document.xml" not in names:
                raise DocumentParseError("DOCX_INVALID", "文件内容不是有效的 DOCX")
            total_size = sum(item.file_size for item in archive.infolist())
            if total_size > max_uncompressed_bytes:
                raise DocumentParseError(
                    "DOCX_ARCHIVE_TOO_LARGE",
                    "DOCX 解压后的内容过大，请拆分文件后上传",
                )
    except DocumentParseError:
        raise
    except BadZipFile as exc:
        raise DocumentParseError("DOCX_INVALID", "DOCX 已损坏或无法解析") from exc


def _normalize_extracted_text(text: str) -> str:
    lines = [re.sub(r"[\t ]+", " ", line).strip() for line in text.splitlines()]
    compacted: list[str] = []
    previous_blank = False
    for line in lines:
        if line:
            compacted.append(line)
            previous_blank = False
        elif compacted and not previous_blank:
            compacted.append("")
            previous_blank = True
    return "\n".join(compacted).strip()
